import openpyxl
from openpyxl.utils import get_column_letter


class ExcelHandler:
    def __init__(self):
        self.path = None
        self.wb = None
        self.ws = None
        self.headers = []          # список заголовков
        self.data_rows = []

    def load(self, path):
        self.path = path
        self.wb = openpyxl.load_workbook(path, data_only=True)
        self.ws = self.wb.active
        self._parse()

    def _parse(self):
        rows = list(self.ws.iter_rows(values_only=True))
        if not rows:
            self.headers = []
            self.data_rows = []
            return

        first_row = rows[0]
        self.headers = [
            (str(c).strip() if c is not None else "") for c in first_row
        ]

        self.data_rows = []
        for i, row in enumerate(rows[1:], start=2):
            values = list(row)
            # достраиваем строку до длины заголовков
            while len(values) < len(self.headers):
                values.append(None)
            self.data_rows.append((i, values))

    def is_loaded(self):
        return self.ws is not None

    def column_count(self):
        return len(self.headers)

    def get_columns_info(self):
        """
        возвращает список словарей с описанием каждого столбца
        индекс, буква столбца, заголовок, пример значения
        используется для выпадающих списков выбора столбца в интерфейсе
        """
        info = []
        sample_row = self.data_rows[0][1] if self.data_rows else []
        for idx in range(len(self.headers)):
            letter = get_column_letter(idx + 1)
            header = self.headers[idx]
            sample = ""
            if idx < len(sample_row) and sample_row[idx] is not None:
                sample = str(sample_row[idx])
                if len(sample) > 45:
                    sample = sample[:45] + "…"
            label = f"{letter}"
            if header:
                label += f" — {header}"
            if sample:
                label += f"  ({sample})"
            info.append(
                {
                    "index": idx,
                    "letter": letter,
                    "header": header,
                    "sample": sample,
                    "label": label,
                }
            )
        return info

    def get_data_rows(self):
        """список кортежей (номер_строки_в_файле, [значения])"""
        return self.data_rows

    def cell_value(self, row_values, col_index):
        if col_index is None:
            return None
        if 0 <= col_index < len(row_values):
            return row_values[col_index]
        return None

    def delete_rows_by_sheet_index(self, sheet_row_indices):
        """удаляет строки из файла (по номерам строк листа) и сохраняет файл"""
        for r in sorted(set(sheet_row_indices), reverse=True):
            self.ws.delete_rows(r)
        self.wb.save(self.path)
        self._parse()

    def save(self):
        if self.wb and self.path:
            self.wb.save(self.path)
